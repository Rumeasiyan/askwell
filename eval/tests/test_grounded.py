"""`eval.grounded`'s own logic — citation matching and the fixture corpus
generator's output — without a database or a model.

Seeding the corpus and driving a real `ask` turn need real Postgres and a
running native inference process; that half is exercised by hand per
`eval/suites/grounded_qa.v1.json`'s own testing notes, not here (no network,
no database — same rule as every other unmarked test, `AGENTS.md` §6).
"""

from pathlib import Path

import docx
import openpyxl
from eval.fixtures.generate_corpus import (
    FIGURES_ROWS,
    HANDBOOK_A_PAGES,
    HANDBOOK_B_PAGES,
    NOTICE_SCAN_LINES,
    SPEC_SECTIONS,
    build_handbook_a,
    build_handbook_b,
    build_notice_scan,
)
from eval.grounded import FIXTURES_DIR, _citation_score
from eval.suite import Task


def _task(**overrides: object) -> Task:
    base = dict(
        id="t",
        prompt="hi",
        scorer="contains_all",
        expected="hi",
        timeout_seconds=1.0,
        expected_documents=("handbook_a.pdf",),
        expected_passages=("eleven paid holiday days",),
    )
    base.update(overrides)
    return Task(**base)  # type: ignore[arg-type]


def _citation(filename: str, passage: str) -> dict[str, object]:
    return {"filename": filename, "passage": passage}


def test_citation_score_matches_document_and_passage() -> None:
    task = _task()
    citations = [
        _citation("handbook_a.pdf", "Meridian Loom employees accrue eleven paid holiday days.")
    ]
    assert _citation_score(task, citations) == 1.0


def test_citation_score_is_case_insensitive() -> None:
    task = _task()
    citations = [_citation("handbook_a.pdf", "ELEVEN PAID HOLIDAY DAYS accrue each year.")]
    assert _citation_score(task, citations) == 1.0


def test_citation_score_rejects_right_passage_wrong_document() -> None:
    task = _task()
    citations = [_citation("handbook_b.pdf", "eleven paid holiday days")]
    assert _citation_score(task, citations) == 0.0


def test_citation_score_rejects_right_document_wrong_passage() -> None:
    task = _task()
    citations = [_citation("handbook_a.pdf", "sixty-three days notice")]
    assert _citation_score(task, citations) == 0.0


def test_citation_score_accepts_either_of_two_expected_documents() -> None:
    """The "answer appears in two places" edge case: a duplicated fact in two
    fixture documents, either citation counts."""
    task = _task(expected_documents=("handbook_a.pdf", "notice_scan.pdf"))
    citations = [_citation("notice_scan.pdf", "eleven paid holiday days")]
    assert _citation_score(task, citations) == 1.0


def test_citation_score_with_no_citations_is_zero() -> None:
    assert _citation_score(_task(), []) == 0.0


def test_fixture_corpus_is_committed_and_reproducible() -> None:
    """The committed bytes under `eval/fixtures/corpus/` match a fresh build
    from the same fact strings — the ticket's own "fixture corpus is
    committed and reproducible" acceptance criterion.

    Byte-for-byte for the PDFs, built from scratch with no embedded
    timestamp. `.docx`/`.xlsx` are zip containers python-docx/openpyxl stamp
    with the current time on every save, so those two are compared by
    content instead — still reproducible, just not byte-identical.
    """
    assert (FIXTURES_DIR / "handbook_a.pdf").read_bytes() == build_handbook_a()
    assert (FIXTURES_DIR / "handbook_b.pdf").read_bytes() == build_handbook_b()
    assert (FIXTURES_DIR / "notice_scan.pdf").read_bytes() == build_notice_scan()

    document = docx.Document(str(FIXTURES_DIR / "spec.docx"))
    facts = [p.text for p in document.paragraphs if p.text in dict(SPEC_SECTIONS).values()]
    assert facts == [fact for _heading, fact in SPEC_SECTIONS]

    workbook = openpyxl.load_workbook(FIXTURES_DIR / "figures.xlsx")
    sheet = workbook.active
    rows = [tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert rows == FIGURES_ROWS


def test_fixture_corpus_covers_the_ticket_scope() -> None:
    """Digital PDFs, a scan, an Office document and a table — the ticket's
    own scope list, `M2-EVAL-TEST-064`."""
    names = {p.name for p in Path(FIXTURES_DIR).iterdir()}
    assert names == {
        "handbook_a.pdf",
        "handbook_b.pdf",
        "notice_scan.pdf",
        "spec.docx",
        "figures.xlsx",
    }
    assert len(HANDBOOK_A_PAGES) == 8
    assert len(HANDBOOK_B_PAGES) == 8
    assert len(NOTICE_SCAN_LINES) == 5
    assert len(SPEC_SECTIONS) == 6
    assert len(FIGURES_ROWS) == 5
