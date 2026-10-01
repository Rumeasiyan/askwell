"""`askwell.sheet_citations` against a real workbook: loaded into its
folder's sandbox database, its rows written as the document's anchors and
passages, and a query's result mapped back to them. `M11-FIX-BE-231`.

The last test drives `askwell.ask._run_sql_turn` and `_finish_sql_turn` over
the same workbook with only the model call faked, so the citation is shown
to reach the `citation` event and the `citations` table from a real query,
not from a hand-built result (#896).
"""

import hashlib
import io
import os
import uuid
from collections.abc import AsyncIterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from askwell import ask as ask_module
from askwell.config import Settings
from askwell.db.engine import session_scope
from askwell.extract_xlsx import ANCHOR_KIND, _read
from askwell.filetypes import WORKBOOK_MIME
from askwell.sheet_citations import SheetCitation, cite_sheet_result
from askwell.sql_execute import execute_sandbox_query
from askwell.table_load import load_workbook_tables

pytestmark = pytest.mark.requires_db

TABLES = (
    "sources, documents, chunks, citations, conversations, messages, clarifications, "
    "memory, schema_notes, reapply_jobs, audit_decisions, audit_interactions, settings"
)

_FIGURES: list[list[object]] = [
    ["Department", "Q1 Revenue", "Headcount", "Avg Tenure Years"],
    ["Textiles", 482000, 34, 4.1],
    ["Logistics", 215000, 19, 2.7],
    ["Research", 903000, 27, 5.6],
    ["Design", 76000, 19, 3.3],
]


@pytest_asyncio.fixture
async def factory(database_url: str) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    async_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)
    engine = create_async_engine(async_url)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as opened:
        await opened.execute(text(f"TRUNCATE {TABLES} CASCADE"))
        await opened.commit()
    yield sessions
    async with sessions() as opened:
        await opened.execute(text(f"TRUNCATE {TABLES} CASCADE"))
        await opened.commit()
    await engine.dispose()


@pytest.fixture
def sheet_settings(sandbox_admin_url: str, tmp_path: Path) -> Settings:
    return Settings(
        database_url="postgresql://x:x@127.0.0.1:1/askwell",  # type: ignore[arg-type]
        sandbox_database_url=sandbox_admin_url,  # type: ignore[arg-type]
        sandbox_owner_password=os.environ["TEST_SANDBOX_OWNER_PASSWORD"],  # type: ignore[arg-type]
        sandbox_readonly_password=os.environ["TEST_SANDBOX_READONLY_PASSWORD"],  # type: ignore[arg-type]
        sql_statement_timeout_seconds=5,
        trace_dir=tmp_path / "traces",
    )


@dataclass(frozen=True)
class _Workbook:
    source_id: uuid.UUID
    document_id: uuid.UUID
    database: str


def _write_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = openpyxl.Workbook()
    first = True
    for title, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        assert sheet is not None
        sheet.title = title
        for row in rows:
            sheet.append(row)
        first = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buffer.getvalue())


async def _index(
    factory: async_sessionmaker[AsyncSession], document_id: uuid.UUID, path: Path
) -> uuid.UUID:
    """Write the workbook's anchors as `extract_xlsx` does and one live
    passage over all of them, as `chunk` packs a small workbook. Returns the
    passage's id."""
    rows, _, _ = _read(str(path))
    chunk_id = uuid.uuid4()
    async with session_scope(factory) as session:
        await session.execute(
            text("DELETE FROM document_pages WHERE document_id = :id"), {"id": document_id}
        )
        await session.execute(
            text("UPDATE documents SET page_count = :count, anchor_kind = :kind WHERE id = :id"),
            {"count": len(rows), "kind": ANCHOR_KIND, "id": document_id},
        )
        for number, (label, row_text) in enumerate(rows, start=1):
            await session.execute(
                text(
                    "INSERT INTO document_pages "
                    "(document_id, page_number, text, has_text, anchor_label) "
                    "VALUES (:id, :number, :text, true, :label)"
                ),
                {"id": document_id, "number": number, "text": row_text, "label": label},
            )
        await session.execute(
            text(
                "UPDATE chunks SET superseded_at = now(), embedding = NULL, content_tsv = NULL "
                "WHERE document_id = :id AND superseded_at IS NULL"
            ),
            {"id": document_id},
        )
        await session.execute(
            text(
                "INSERT INTO chunks (id, document_id, ordinal, page_from, page_to, content) "
                "VALUES (:id, :document_id, 0, 1, :page_to, :content)"
            ),
            {
                "id": chunk_id,
                "document_id": document_id,
                "page_to": len(rows),
                "content": "\n\n".join(row_text for _, row_text in rows),
            },
        )
    return chunk_id


async def _load(
    factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    root: Path,
    sheets: dict[str, list[list[object]]],
) -> _Workbook:
    path = root / "figures.xlsx"
    _write_workbook(path, sheets)
    async with session_scope(factory) as session:
        source_id = (
            await session.execute(
                text(
                    "INSERT INTO sources (kind, name, root_path, status) "
                    "VALUES ('file', :name, :root, 'ready') RETURNING id"
                ),
                {"name": root.name, "root": str(root)},
            )
        ).scalar_one()
        document_id = (
            await session.execute(
                text(
                    "INSERT INTO documents "
                    "(source_id, filename, path, mime, sha256, version, status) "
                    "VALUES (:source_id, 'figures.xlsx', :path, :mime, :sha256, 1, 'ready') "
                    "RETURNING id"
                ),
                {
                    "source_id": source_id,
                    "path": str(path),
                    "mime": WORKBOOK_MIME,
                    "sha256": hashlib.sha256(os.urandom(16)).hexdigest(),
                },
            )
        ).scalar_one()
    source_id, document_id = uuid.UUID(str(source_id)), uuid.UUID(str(document_id))
    await _index(factory, document_id, path)
    outcome = await load_workbook_tables(factory, settings, source_id, document_id)
    assert outcome.failure is None and outcome.tables
    async with session_scope(factory) as session:
        database = (
            await session.execute(
                text("SELECT sandbox_db FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).scalar_one()
    return _Workbook(source_id, document_id, str(database))


async def _cite(
    factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    workbook: _Workbook,
    query: str,
) -> list[SheetCitation]:
    async with session_scope(factory) as session:
        result = await execute_sandbox_query(
            session, settings, database=workbook.database, query=query
        )
        return await cite_sheet_result(
            session,
            settings,
            source_id=workbook.source_id,
            database=workbook.database,
            query=query,
            rows=result.rows,
            truncated=False,
        )


async def test_a_filtered_row_cites_that_row_of_the_workbook(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics' LIMIT 1000",
    )

    assert len(citations) == 1
    (citation,) = citations
    assert citation.document_id == workbook.document_id
    assert citation.filename == "figures.xlsx"
    assert citation.anchor_kind == ANCHOR_KIND
    assert citation.quoted_span == "Logistics | 215000 | 19 | 2.7"
    assert citation.quoted_span in citation.passage
    assert citation.heading == "Data, row 3"
    # The row's own anchor, so the viewer opens on it.
    assert (citation.page_from, citation.page_to) == (3, 3)


async def test_an_aggregate_over_many_rows_cites_the_sheet_not_every_row(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})

    citations = await _cite(
        factory, sheet_settings, workbook, "SELECT sum(headcount) FROM figures_xlsx_data"
    )

    assert len(citations) == 1
    (citation,) = citations
    assert citation.heading == "Data"
    assert citation.quoted_span is None
    assert "Department | Q1 Revenue | Headcount" in citation.passage


async def test_a_tie_the_limit_cut_cites_the_sheet_never_a_row_not_returned(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    """#896: two departments at 19 and `LIMIT 1`. Which one came back cannot
    be known from a second read, so no row is cited (C4)."""
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_data WHERE headcount < 20 ORDER BY headcount LIMIT 1",
    )

    assert [(c.heading, c.quoted_span) for c in citations] == [("Data", None)]


async def test_two_sheets_cite_the_sheet_the_query_read(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(
        factory,
        sheet_settings,
        tmp_path / "folder",
        {"North": _FIGURES, "South": [_FIGURES[0], ["Retail", 118000, 41, 1.9]]},
    )

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_south WHERE department = 'Retail'",
    )

    assert [(c.heading, c.quoted_span) for c in citations] == [
        ("South, row 2", "Retail | 118000 | 41 | 1.9")
    ]


async def test_a_reindexed_workbook_cites_its_live_passage_never_a_retired_one(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})
    retired = await _live_chunk(factory, workbook.document_id)
    live = await _index(factory, workbook.document_id, tmp_path / "folder" / "figures.xlsx")
    assert live != retired

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics'",
    )

    assert [c.chunk_id for c in citations] == [live]


async def test_a_row_its_passages_no_longer_hold_is_not_cited(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    """The workbook was edited and re-indexed since its tables loaded: the
    live passage no longer says what the table row does, so neither the row
    nor the sheet is cited from it."""
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})
    path = tmp_path / "folder" / "figures.xlsx"
    _write_workbook(path, {"Data": [_FIGURES[0], ["Logistics", 215000, 23, 2.7]]})
    await _index(factory, workbook.document_id, path)

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics'",
    )

    assert citations == []


async def test_a_sheet_with_no_passage_left_is_not_cited(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})
    async with session_scope(factory) as session:
        await session.execute(
            text(
                "UPDATE chunks SET superseded_at = now(), embedding = NULL, content_tsv = NULL "
                "WHERE document_id = :id"
            ),
            {"id": workbook.document_id},
        )

    for query in (
        "SELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics'",
        "SELECT sum(headcount) FROM figures_xlsx_data",
    ):
        assert await _cite(factory, sheet_settings, workbook, query) == []


async def test_a_deleted_workbooks_tables_cite_nothing(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})
    async with session_scope(factory) as session:
        await session.execute(
            text("UPDATE documents SET deleted_at = now() WHERE id = :id"),
            {"id": workbook.document_id},
        )

    citations = await _cite(
        factory,
        sheet_settings,
        workbook,
        "SELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics'",
    )

    assert citations == []


async def _live_chunk(factory: async_sessionmaker[AsyncSession], document_id: uuid.UUID) -> Any:
    async with session_scope(factory) as session:
        return (
            await session.execute(
                text("SELECT id FROM chunks WHERE document_id = :id AND superseded_at IS NULL"),
                {"id": document_id},
            )
        ).scalar_one()


# --- a real SQL turn, model call faked -------------------------------------


@dataclass
class _Completion:
    text: str


@dataclass
class _FakeInferenceClient:
    text: str

    async def generate(
        self, _prompt: str, *, max_tokens: int = 512, temperature: float = 0.2
    ) -> _Completion:
        return _Completion(text=self.text)


async def test_a_sql_turn_over_a_sheet_streams_and_stores_its_workbook_citation(
    factory: async_sessionmaker[AsyncSession], sheet_settings: Settings, tmp_path: Path
) -> None:
    workbook = await _load(factory, sheet_settings, tmp_path / "folder", {"Data": _FIGURES})
    client = _FakeInferenceClient(
        "```sql\nSELECT headcount FROM figures_xlsx_data WHERE department = 'Logistics'\n```"
    )
    question = "How many people work in the Logistics department, per the department figures?"

    async with session_scope(factory) as session:
        sql_answer = await ask_module._run_sql_turn(
            sheet_settings, session, client, question=question, source_id=None
        )
    assert sql_answer is not None and sql_answer.sql_result is not None
    assert sql_answer.text == "Found 1 row: headcount 19."
    assert sql_answer.trace_step["citations"] == 1

    conversation_id, message_id = uuid.uuid4(), uuid.uuid4()
    async with session_scope(factory) as session:
        await session.execute(
            text("INSERT INTO conversations (id) VALUES (:id)"), {"id": conversation_id}
        )
    turn = ask_module._Turn(message_id=message_id, conversation_id=conversation_id)
    await ask_module._finish_sql_turn(
        sheet_settings,
        factory,
        turn,
        question=question,
        source_id=None,
        sql_answer=sql_answer,
        model_name="test-model",
        turn_started=0.0,
    )

    events = [event.data for event in turn.events if event.kind == "citation"]
    assert len(events) == 1
    assert events[0]["filename"] == "figures.xlsx"
    assert events[0]["document_id"] == str(workbook.document_id)
    assert "Logistics | 215000 | 19" in events[0]["passage"]
    assert events[0]["quoted_span"] == "Logistics | 215000 | 19 | 2.7"
    done = next(event.data for event in turn.events if event.kind == "done")
    assert done["status"] == "completed"
    assert done["source_count"] == 1
    async with session_scope(factory) as session:
        stored = (
            await session.execute(
                text("SELECT chunk_id, quoted_span FROM citations WHERE message_id = :id"),
                {"id": message_id},
            )
        ).all()
    assert [(str(chunk_id), span) for chunk_id, span in stored] == [
        (events[0]["chunk_id"], "Logistics | 215000 | 19 | 2.7")
    ]
