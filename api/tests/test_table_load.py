"""Loading an inferred table into the sandbox. `M4-CSV-ING-094`.

Pure-function tests only — identifier normalisation, type/cast decisions, and
date parsing, all against literals with no database. Everything that depends
on the real sandbox instance (`_load_inferences` end to end, the reload path,
the caps) is `requires_db` and lives in `test_table_load_db.py`, the same
split `test_table_infer.py`/`test_table_infer_db.py` and
`test_dump_import.py` already use.
"""

import io
from datetime import date

import openpyxl
import pytest

from askwell.table_infer import (
    ColumnInference,
    DateFormatDetection,
    DateFormatVerdict,
    infer_csv,
    infer_xlsx,
)
from askwell.table_load import (
    RowFailure,
    cast_value,
    normalise_columns,
    normalise_identifier,
    sheet_skip_reason,
    sql_type_for,
    workbook_comment,
    workbook_table_prefix,
)

# --- identifier normalisation -------------------------------------------------


def test_a_valid_identifier_is_only_lowercased() -> None:
    assert normalise_identifier("amount", set()) == "amount"


def test_spaces_and_punctuation_become_underscores() -> None:
    assert normalise_identifier("Reference #", set()) == "reference"
    assert normalise_identifier("Unit Price", set()) == "unit_price"


def test_a_name_starting_with_a_digit_is_prefixed() -> None:
    assert normalise_identifier("2026 total", set()) == "c_2026_total"


def test_an_empty_or_all_punctuation_name_falls_back_to_column() -> None:
    assert normalise_identifier("!!!", set()) == "column"


def test_two_names_that_normalise_the_same_way_are_deduplicated() -> None:
    seen: set[str] = set()
    first = normalise_identifier("Amount", seen)
    second = normalise_identifier("amount!", seen)
    assert first == "amount"
    assert second == "amount_2"


def test_normalise_columns_applies_one_shared_seen_set_in_order() -> None:
    assert normalise_columns(["Name", "name", "Name "]) == ["name", "name_2", "name_3"]


# --- sql_type_for --------------------------------------------------------------


def _column(
    inferred_type: str = "string",
    *,
    ambiguous: bool = False,
    date_format: DateFormatDetection | None = None,
) -> ColumnInference:
    return ColumnInference(
        name="col",
        inferred_type=inferred_type,
        confidence=1.0,
        sample_values=["x"],
        ambiguous=ambiguous,
        date_format=date_format,
    )


def test_an_ambiguous_column_always_loads_as_text() -> None:
    column = _column("decimal", ambiguous=True)
    assert sql_type_for(column, day_first_override=None) == ("text", None)


def test_a_confident_integer_column_maps_to_bigint() -> None:
    column = _column("integer")
    assert sql_type_for(column, day_first_override=None) == ("bigint", None)


def test_a_confident_decimal_column_maps_to_numeric() -> None:
    column = _column("decimal")
    assert sql_type_for(column, day_first_override=None) == ("numeric", None)


def test_a_confident_boolean_column_maps_to_boolean() -> None:
    column = _column("boolean")
    assert sql_type_for(column, day_first_override=None) == ("boolean", None)


def test_a_disambiguated_day_first_date_column_loads_as_date() -> None:
    column = _column(
        "date", date_format=DateFormatDetection(DateFormatVerdict.DAY_FIRST, "13/01/2026", "...")
    )
    assert sql_type_for(column, day_first_override=None) == ("date", True)


def test_a_disambiguated_month_first_date_column_loads_as_date() -> None:
    column = _column(
        "date",
        date_format=DateFormatDetection(DateFormatVerdict.MONTH_FIRST, "01/13/2026", "..."),
    )
    assert sql_type_for(column, day_first_override=None) == ("date", False)


def test_an_iso_only_date_column_needs_no_day_first_decision() -> None:
    column = _column("date", date_format=DateFormatDetection(DateFormatVerdict.NOT_APPLICABLE))
    assert sql_type_for(column, day_first_override=None) == ("date", None)


def test_an_undecided_date_column_loads_as_text_until_an_override_resolves_it() -> None:
    column = _column(
        "date",
        ambiguous=True,
        date_format=DateFormatDetection(DateFormatVerdict.AMBIGUOUS),
    )
    assert sql_type_for(column, day_first_override=None) == ("text", None)
    assert sql_type_for(column, day_first_override=True) == ("date", True)
    assert sql_type_for(column, day_first_override=False) == ("date", False)


# --- cast_value ------------------------------------------------------------


def test_a_blank_cell_casts_to_none_regardless_of_type() -> None:
    assert cast_value("  ", "bigint", None) is None
    assert cast_value("", "text", None) is None


def test_bigint_and_numeric_cast() -> None:
    assert cast_value("42", "bigint", None) == 42
    assert cast_value("1,200.50", "numeric", None) == 1200.5


def test_an_unparseable_integer_raises_value_error() -> None:
    with pytest.raises(ValueError, match=r"not a number|invalid literal"):
        cast_value("abc", "bigint", None)


def test_boolean_recognises_the_documented_spellings() -> None:
    assert cast_value("Yes", "boolean", None) is True
    assert cast_value("0", "boolean", None) is False


def test_an_unrecognised_boolean_raises() -> None:
    with pytest.raises(ValueError, match="not a recognised boolean"):
        cast_value("maybe", "boolean", None)


def test_iso_date_casts_without_a_day_first_decision() -> None:
    assert cast_value("2026-01-13", "date", None) == date(2026, 1, 13)


def test_named_month_date_casts() -> None:
    assert cast_value("4 January 2026", "date", None) == date(2026, 1, 4)


def test_numeric_date_casts_day_first() -> None:
    assert cast_value("13/01/2026", "date", True) == date(2026, 1, 13)


def test_numeric_date_casts_month_first() -> None:
    assert cast_value("01/13/2026", "date", False) == date(2026, 1, 13)


def test_a_numeric_date_with_no_day_first_decision_raises() -> None:
    with pytest.raises(ValueError, match="day/month order was never decided"):
        cast_value("03/04/2026", "date", None)


def test_text_cast_returns_the_value_unchanged() -> None:
    assert cast_value("  Anna  ", "text", None) == "Anna"


# --- row numbering, shared vocabulary with table_infer ------------------------


def test_row_failure_is_a_plain_row_number_and_reason() -> None:
    failure = RowFailure(3, "'abc' is not a number")
    assert failure.row_number == 3
    assert "not a number" in failure.reason


def test_infer_csv_carries_the_data_rows_for_loading_to_use() -> None:
    inference = infer_csv("people.csv", b"name,amount\nAnna,10\nBen,20\n")
    assert inference.rows == [["Anna", "10"], ["Ben", "20"]]


# --- workbook sheets (`M11-FIX-ING-224`) ----------------------------------------


def _workbook(sheets: dict[str, list[list[object]]], *, merge: str | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    first = True
    for title, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        assert sheet is not None
        sheet.title = title
        for row in rows:
            sheet.append(row)
        if merge and first:
            sheet.merge_cells(merge)
        first = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_FIGURES = [
    ["Department", "Q1 Revenue", "Headcount", "Avg Tenure Years"],
    ["Textiles", 482000, 34, 4.1],
    ["Logistics", 215000, 19, 2.7],
]


def test_a_sheet_with_a_header_row_over_typed_data_is_loadable() -> None:
    [sheet] = infer_xlsx("figures.xlsx", _workbook({"Figures": _FIGURES}))
    assert sheet_skip_reason(sheet) is None


def test_a_sheet_whose_first_row_is_data_is_skipped_with_the_reason() -> None:
    [sheet] = infer_xlsx("figures.xlsx", _workbook({"Raw": [[1, 2], [3, 4], [5, 6]]}))
    reason = sheet_skip_reason(sheet)
    assert reason is not None
    assert "header" in reason


def test_a_sheet_of_free_text_is_skipped_because_its_header_cannot_be_told_apart() -> None:
    [sheet] = infer_xlsx("notes.xlsx", _workbook({"Notes": [["alpha"], ["beta"], ["gamma"]]}))
    reason = sheet_skip_reason(sheet)
    assert reason is not None
    assert "header" in reason


def test_an_empty_sheet_is_skipped() -> None:
    [sheet] = infer_xlsx("empty.xlsx", _workbook({"Blank": []}))
    assert sheet_skip_reason(sheet) == "the sheet is empty"


def test_a_merged_header_is_skipped_rather_than_guessed_at() -> None:
    raw = _workbook(
        {"Grouped": [["group", "", "amount"], ["Anna", "x", 10], ["Ben", "y", 20]]}, merge="A1:B1"
    )
    [sheet] = infer_xlsx("grouped.xlsx", raw)
    reason = sheet_skip_reason(sheet)
    assert reason is not None
    assert "merged" in reason


def test_two_sheets_with_the_same_headers_are_both_loadable_under_their_own_names() -> None:
    raw = _workbook({"North": _FIGURES, "South": _FIGURES})
    north, south = infer_xlsx(workbook_table_prefix("sub/figures.xlsx")[:-1], raw)
    assert (north.table_name, south.table_name) == (
        "sub/figures.xlsx:North",
        "sub/figures.xlsx:South",
    )
    assert sheet_skip_reason(north) is None
    assert sheet_skip_reason(south) is None


def test_a_workbooks_notes_and_tables_are_keyed_on_its_path_inside_the_source() -> None:
    assert workbook_table_prefix("sub/figures.xlsx") == "sub/figures.xlsx:"
    assert workbook_comment("sub/figures.xlsx") == "askwell workbook: sub/figures.xlsx"


def test_a_sheet_name_that_is_not_an_identifier_normalises_to_one() -> None:
    assert (
        normalise_identifier("figures.xlsx:Quarterly Department Figures", set())
        == "figures_xlsx_quarterly_department_figures"
    )
