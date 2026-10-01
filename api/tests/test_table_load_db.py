"""Loading and reloading a table into the real sandbox instance.
`M4-CSV-ING-094`. Same split as `test_dump_import.py`: pure-Python casting
and identifier logic is in `test_table_load.py`; everything here needs a
real, separate Postgres instance to prove.
"""

import hashlib
import io
import json
import os
import uuid
from collections.abc import AsyncIterator
from pathlib import Path

import openpyxl
import psycopg
import pytest
import pytest_asyncio
from pydantic import SecretStr
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from askwell import ingest
from askwell.config import Settings
from askwell.dump_import import set_dump_size_cap_bytes
from askwell.filetypes import WORKBOOK_MIME
from askwell.review import answer_clarification
from askwell.table_load import (
    WORKBOOK_TABLES_FAILED,
    WORKBOOK_TABLES_LOADED,
    EmptyTable,
    TableCapExceeded,
    create_table_source,
    load_workbook_tables,
    process_table_source,
    reload_source,
    sweep_workbook_tables,
)

pytestmark = pytest.mark.requires_db

TABLES = (
    "sources, documents, clarifications, memory, schema_notes, reapply_jobs, "
    "audit_decisions, settings"
)


@pytest_asyncio.fixture
async def factory(database_url: str) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    async_url = database_url.replace("postgresql://", "postgresql+psycopg://", 1)
    engine = create_async_engine(async_url)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as opened:
        await opened.execute(text(f"TRUNCATE {TABLES} CASCADE"))
        await opened.commit()
    yield sessions
    async with sessions() as opened:
        await opened.execute(text(f"TRUNCATE {TABLES} CASCADE"))
        await opened.commit()
    await engine.dispose()


@pytest.fixture
def table_settings(sandbox_admin_url: str) -> Settings:
    return Settings(
        database_url="postgresql://x:x@127.0.0.1:1/askwell",  # type: ignore[arg-type]
        sandbox_database_url=sandbox_admin_url,  # type: ignore[arg-type]
        sandbox_owner_password=os.environ["TEST_SANDBOX_OWNER_PASSWORD"],  # type: ignore[arg-type]
        sandbox_readonly_password=os.environ["TEST_SANDBOX_READONLY_PASSWORD"],  # type: ignore[arg-type]
    )


async def _make_source(
    factory: async_sessionmaker[AsyncSession], name: str, csv_path: Path
) -> uuid.UUID:
    async with factory() as session:
        source_id = await create_table_source(session, name, str(csv_path))
        await session.commit()
    return source_id


async def test_a_clean_csv_loads_as_a_real_table_with_confirmed_types(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    csv_path = tmp_path / "people.csv"
    csv_path.write_text("Name,Amount\nAnna,10\nBen,20\n")
    source_id = await _make_source(factory, "people", csv_path)

    results = await process_table_source(factory, table_settings, source_id, str(csv_path))

    assert len(results) == 1
    result = results[0]
    assert result.row_count == 2
    assert result.failed_rows == []

    async with factory() as session:
        row = (
            await session.execute(
                text("SELECT status, sandbox_db FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).one()
        assert row[0] == "ready"
        database = row[1]
        assert database is not None

        notes = (
            await session.execute(
                text(
                    "SELECT column_name FROM schema_notes WHERE source_id = :id "
                    "AND origin = 'inferred'"
                ),
                {"id": source_id},
            )
        ).all()
        # `M4-CSV-ING-092`'s own per-column notes, plus this ticket's
        # table-level note (`column_name IS NULL`).
        assert None in {row[0] for row in notes}

    admin_url = table_settings.sandbox_database_url.get_secret_value()
    import psycopg

    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        rows = conn.execute("SELECT name, amount FROM people_csv ORDER BY name").fetchall()
    assert rows == [("Anna", 10), ("Ben", 20)]


async def test_a_column_name_that_is_not_a_valid_identifier_is_normalised_and_recorded(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    csv_path = tmp_path / "export.csv"
    csv_path.write_text("Reference #,Amount\nA1,10\nA2,20\n")
    source_id = await _make_source(factory, "export", csv_path)

    results = await process_table_source(factory, table_settings, source_id, str(csv_path))
    assert results[0].sql_table_name == "export_csv"

    async with factory() as session:
        note = (
            await session.execute(
                text(
                    "SELECT description FROM schema_notes WHERE source_id = :id "
                    "AND column_name = 'Reference #' AND description LIKE 'Loaded as column%'"
                ),
                {"id": source_id},
            )
        ).first()
        assert note is not None
        assert "reference" in note[0]
        assert "Reference #" in note[0]


async def test_malformed_rows_are_reported_by_row_number_not_dropped_silently(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    # Nine good integers and one `oops` — confident enough (9/10 = 90%,
    # over `_MIN_TYPE_CONFIDENCE`) that `table_infer` still calls the column
    # `integer` rather than falling back to an ambiguous `string`, so the
    # bad row only ever surfaces as a real cast failure during the load.
    good_rows = "\n".join(f"person{i},{i * 10}" for i in range(1, 10))
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text(f"name,amount\n{good_rows}\nBen,oops\n")
    source_id = await _make_source(factory, "sales", csv_path)

    results = await process_table_source(factory, table_settings, source_id, str(csv_path))

    result = results[0]
    assert result.row_count == 9
    assert len(result.failed_rows) == 1
    # Row 1 is the header, so the 10th data row ("Ben,oops") is row 11.
    assert result.failed_rows[0].row_number == 11
    assert "oops" in result.failed_rows[0].reason


async def test_a_100k_row_csv_loads_under_the_default_caps(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """`M11-FIX-ING-234`, #862: the CSV counterpart of the 100k-row sheet,
    through the same `_load_table`, under caps nobody changed."""
    csv_path = tmp_path / "ledger.csv"
    csv_path.write_text("name,amount\n" + "".join(f"row{i},{i}\n" for i in range(100_000)))
    source_id = await _make_source(factory, "ledger", csv_path)

    [result] = await process_table_source(factory, table_settings, source_id, str(csv_path))

    assert result.row_count == 100_000
    assert result.failed_rows == []
    status, database = await _source_row(factory, source_id)
    assert status == "ready"
    assert database is not None
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        counted = conn.execute("SELECT count(*), sum(amount) FROM ledger_csv").fetchone()
    assert counted == (100_000, sum(range(100_000)))


async def test_a_row_the_database_refuses_is_reported_by_row_number_and_its_batch_still_loads(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """`M11-FIX-ING-234`'s edge case. A value Python casts but Postgres
    refuses (an integer past `bigint`) fails its whole batch, which is then
    retried row by row: that row alone is reported, with its number, and
    every other row of the batch is loaded. A Python cast failure in the same
    batch is reported alongside it, in row order."""
    rows = [f"row{i},{i}" for i in range(2_500)]
    rows[1_200] = "row1200,99999999999999999999"  # row 1202, counting the header
    rows[1_500] = "row1500,oops"  # row 1502
    csv_path = tmp_path / "ledger.csv"
    csv_path.write_text("name,amount\n" + "\n".join(rows) + "\n")
    source_id = await _make_source(factory, "ledger", csv_path)

    [result] = await process_table_source(factory, table_settings, source_id, str(csv_path))

    assert [failure.row_number for failure in result.failed_rows] == [1_202, 1_502]
    assert "out of range" in result.failed_rows[0].reason
    assert "oops" in result.failed_rows[1].reason
    assert result.row_count == 2_498
    _, database = await _source_row(factory, source_id)
    assert database is not None
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        names = {name for (name,) in conn.execute("SELECT name FROM ledger_csv").fetchall()}
    assert len(names) == 2_498
    assert {"row1199", "row1201", "row1499", "row1501"} <= names
    assert not {"row1200", "row1500"} & names


async def test_an_empty_file_is_refused_with_the_reason(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text("")
    source_id = await _make_source(factory, "empty", csv_path)

    with pytest.raises(EmptyTable):
        await process_table_source(factory, table_settings, source_id, str(csv_path))

    async with factory() as session:
        row = (
            await session.execute(
                text("SELECT status, sandbox_db FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).one()
        # Never reached the sandbox at all — nothing to create, nothing to drop.
        assert row[1] is None


async def test_the_size_cap_aborts_the_load_and_drops_the_database(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    csv_path = tmp_path / "big.csv"
    csv_path.write_text("name\n" + "\n".join(f"row{i}" for i in range(500)) + "\n")
    source_id = await _make_source(factory, "big", csv_path)

    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024)
        await session.commit()

    with pytest.raises(TableCapExceeded) as exc_info:
        await process_table_source(factory, table_settings, source_id, str(csv_path))
    assert exc_info.value.cap == "size"

    async with factory() as session:
        row = (
            await session.execute(
                text("SELECT status, sandbox_db, last_error FROM sources WHERE id = :id"),
                {"id": source_id},
            )
        ).one()
        assert row[0] == "attention"
        assert row[1] is None
        assert "size cap" in row[2]


async def test_a_batch_that_crosses_the_size_cap_mid_load_aborts_it(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """`M11-FIX-ING-234`'s edge case. The 1 KiB cap above is crossed before a
    single row lands (an empty database is megabytes), so it cannot tell a
    per-batch check from none. Here the cap sits just above an empty
    database's size: the first batches fit, a later one crosses it inside the
    uncommitted transaction, and the load stops on the size cap."""
    from tests.conftest_sandbox import TEST_PREFIX

    admin_url = table_settings.sandbox_database_url.get_secret_value()
    probe = f"{TEST_PREFIX}{uuid.uuid4().hex}"
    with psycopg.connect(admin_url, autocommit=True) as admin:
        admin.execute(f'CREATE DATABASE "{probe}"')
        try:
            row = admin.execute("SELECT pg_database_size(%s)", (probe,)).fetchone()
        finally:
            admin.execute(f'DROP DATABASE "{probe}"')
    assert row is not None
    empty_size = int(row[0])

    # ~130 bytes a row: one 1,000-row batch is ~130 KiB, well under the
    # 512 KiB margin, and 20,000 rows are ~2.6 MB, well over it.
    wide = "x" * 100
    csv_path = tmp_path / "wide.csv"
    csv_path.write_text("name\n" + "".join(f"{wide}{i}\n" for i in range(20_000)))
    source_id = await _make_source(factory, "wide", csv_path)
    async with factory() as session:
        await set_dump_size_cap_bytes(session, empty_size + 512 * 1024)
        await session.commit()

    with pytest.raises(TableCapExceeded) as exc_info:
        await process_table_source(factory, table_settings, source_id, str(csv_path))
    assert exc_info.value.cap == "size"
    status, database = await _source_row(factory, source_id)
    assert status == "attention"
    assert database is None


async def test_answering_a_date_format_clarification_reloads_the_column_as_a_real_date(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    # Both parts of each value are `<= 12`, so neither disambiguates itself
    # (`table_infer.detect_date_format` returns `AMBIGUOUS`) and the column
    # loads as `text` until the clarification it raises is answered.
    csv_path = tmp_path / "registrations.csv"
    csv_path.write_text("name,registered\nAnna,03/04/2026\nBen,05/06/2026\n")
    source_id = await _make_source(factory, "registrations", csv_path)

    await process_table_source(factory, table_settings, source_id, str(csv_path))

    async with factory() as session:
        database = (
            await session.execute(
                text("SELECT sandbox_db FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).scalar_one()

    import psycopg

    admin_url = table_settings.sandbox_database_url.get_secret_value()
    dsn = admin_url.rsplit("/", 1)[0] + f"/{database}"

    # Before the clarification is answered, an undecided date column loads
    # as text, verbatim — never guessed.
    with psycopg.connect(dsn, autocommit=True) as conn:
        rows = conn.execute(
            "SELECT data_type FROM information_schema.columns "
            "WHERE table_name = 'registrations_csv' AND column_name = 'registered'"
        ).fetchall()
    assert rows[0][0] == "text"

    async with factory() as session:
        clarification = (
            await session.execute(
                text(
                    "SELECT id, options FROM clarifications WHERE source_id = :id "
                    "AND evidence->>'trigger' = 'date_format'"
                ),
                {"id": source_id},
            )
        ).one()
        clarification_id, options = clarification
        await answer_clarification(session, clarification_id, options[0])
        await session.commit()

    await reload_source(factory, table_settings, source_id)

    with psycopg.connect(dsn, autocommit=True) as conn:
        column = conn.execute(
            "SELECT data_type FROM information_schema.columns "
            "WHERE table_name = 'registrations_csv' AND column_name = 'registered'"
        ).fetchall()
        values = conn.execute(
            "SELECT name, registered FROM registrations_csv ORDER BY name"
        ).fetchall()
    assert column[0][0] == "date"
    # Day-first: `03/04/2026` -> 3 April, `05/06/2026` -> 5 June.
    assert values[0][1].isoformat() == "2026-04-03"
    assert values[1][1].isoformat() == "2026-06-05"


async def test_date_format_answers_are_read_from_evidence_and_from_the_old_subject_shape(
    factory: async_sessionmaker[AsyncSession], tmp_path: Path
) -> None:
    # #361 moved the table out of the subject and into the evidence. A sheet
    # name contains a colon (`data.xlsx:North`), which the old subject parse
    # split in the wrong place. An answer recorded before the change still
    # carries `"{table}: {column}"` and is still honoured.
    from askwell.table_load import _date_format_overrides

    csv_path = tmp_path / "registrations.csv"
    csv_path.write_text("name\nAnna\n")
    source_id = await _make_source(factory, "registrations", csv_path)
    options = json.dumps(["DD/MM/YYYY (day first)", "MM/DD/YYYY (month first)"])
    async with factory() as session:
        for subject, evidence, answer in (
            (
                "registered",
                {"trigger": "date_format", "table_name": "data.xlsx:North"},
                "DD/MM/YYYY (day first)",
            ),
            ("old.csv: joined", {"trigger": "date_format"}, "MM/DD/YYYY (month first)"),
            ("data.xlsx:South: left", {"trigger": "date_format"}, "DD/MM/YYYY (day first)"),
        ):
            await session.execute(
                text(
                    "INSERT INTO clarifications "
                    "(id, source_id, subject, question, options, evidence, status, answer) "
                    "VALUES (:id, :source_id, :subject, 'which?', CAST(:options AS jsonb), "
                    "CAST(:evidence AS jsonb), 'answered', :answer)"
                ),
                {
                    "id": uuid.uuid4(),
                    "source_id": source_id,
                    "subject": subject,
                    "options": options,
                    "evidence": json.dumps(evidence),
                    "answer": answer,
                },
            )
        await session.commit()

        overrides = await _date_format_overrides(session, source_id)

    assert overrides == {
        ("data.xlsx:North", "registered"): True,
        ("old.csv", "joined"): False,
        ("data.xlsx:South", "left"): True,
    }


async def test_the_reapply_job_for_a_date_format_answer_reloads_the_column_itself(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    # #361: the compound subject matched no schema note, so answering queued
    # no job and nothing reloaded unless `reload_source` was called by hand.
    from askwell import reapply

    csv_path = tmp_path / "registrations.csv"
    csv_path.write_text("name,registered\nAnna,03/04/2026\nBen,05/06/2026\n")
    source_id = await _make_source(factory, "registrations", csv_path)
    await process_table_source(factory, table_settings, source_id, str(csv_path))

    async with factory() as session:
        clarification_id, options, database = (
            await session.execute(
                text(
                    "SELECT c.id, c.options, s.sandbox_db FROM clarifications c "
                    "JOIN sources s ON s.id = c.source_id WHERE c.source_id = :id "
                    "AND c.evidence->>'trigger' = 'date_format'"
                ),
                {"id": source_id},
            )
        ).one()
        outcome = await answer_clarification(session, clarification_id, options[0])
        await session.commit()
    assert outcome.reapply_job_id is not None

    await reapply.run_job(factory, table_settings, outcome.reapply_job_id)

    import psycopg

    admin_url = table_settings.sandbox_database_url.get_secret_value()
    dsn = admin_url.rsplit("/", 1)[0] + f"/{database}"
    with psycopg.connect(dsn, autocommit=True) as conn:
        column = conn.execute(
            "SELECT data_type FROM information_schema.columns "
            "WHERE table_name = 'registrations_csv' AND column_name = 'registered'"
        ).fetchall()
    assert column[0][0] == "date"

    async with factory() as session:
        note = (
            await session.execute(
                text(
                    "SELECT origin, description FROM schema_notes WHERE source_id = :id "
                    "AND column_name = 'registered' AND superseded_by IS NULL"
                ),
                {"id": source_id},
            )
        ).one()
    assert note == ("user", options[0])


# --- a workbook's sheets, loaded alongside its document (`M11-FIX-ING-224`) ---


_FIGURES: list[list[object]] = [
    ["Department", "Q1 Revenue", "Headcount", "Avg Tenure Years"],
    ["Textiles", 482000, 34, 4.1],
    ["Logistics", 215000, 19, 2.7],
    ["Research", 903000, 27, 5.6],
]


def _write_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = openpyxl.Workbook()
    first = True
    for title, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        assert sheet is not None
        sheet.title = title
        for row in rows:
            sheet.append(row)
        first = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buffer.getvalue())


async def _folder_source(factory: async_sessionmaker[AsyncSession], root: Path) -> uuid.UUID:
    async with factory() as session:
        source_id = (
            await session.execute(
                text(
                    "INSERT INTO sources (kind, name, root_path, status) "
                    "VALUES ('file', :name, :root, 'ready') RETURNING id"
                ),
                {"name": root.name, "root": str(root)},
            )
        ).scalar_one()
        await session.commit()
    return uuid.UUID(str(source_id))


async def _workbook_document(
    factory: async_sessionmaker[AsyncSession], source_id: uuid.UUID, path: Path
) -> uuid.UUID:
    async with factory() as session:
        document_id = (
            await session.execute(
                text(
                    "INSERT INTO documents "
                    "(source_id, filename, path, mime, sha256, version, status) "
                    "VALUES (:source_id, :filename, :path, :mime, :sha256, 1, 'indexing') "
                    "RETURNING id"
                ),
                {
                    "source_id": source_id,
                    "filename": path.name,
                    "path": str(path),
                    "mime": WORKBOOK_MIME,
                    "sha256": hashlib.sha256(str(path).encode() + os.urandom(8)).hexdigest(),
                },
            )
        ).scalar_one()
        await session.commit()
    return uuid.UUID(str(document_id))


def _sandbox_tables(settings: Settings, database: str) -> dict[str, str | None]:
    admin_url = settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        rows = conn.execute(
            "SELECT c.relname, obj_description(c.oid, 'pg_class') FROM pg_class c "
            "JOIN pg_namespace n ON n.oid = c.relnamespace "
            "WHERE n.nspname = 'public' AND c.relkind = 'r'"
        ).fetchall()
    return {str(name): comment for name, comment in rows}


async def _source_row(
    factory: async_sessionmaker[AsyncSession], source_id: uuid.UUID
) -> tuple[str, str | None]:
    async with factory() as session:
        row = (
            await session.execute(
                text("SELECT status, sandbox_db FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).one()
    return row[0], row[1]


async def _audit(factory: async_sessionmaker[AsyncSession], kind: str) -> list[dict[str, object]]:
    async with factory() as session:
        rows = await session.execute(
            text("SELECT payload FROM audit_decisions WHERE kind = :kind ORDER BY id"),
            {"kind": kind},
        )
        return [dict(row[0]) for row in rows]


async def test_each_headed_sheet_of_a_workbook_loads_as_a_table_in_its_folders_sandbox(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(
        workbook,
        {
            "Quarterly Department Figures": _FIGURES,
            "Raw": [[1, 2], [3, 4], [5, 6]],
            "Blank": [],
        },
    )
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)

    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert outcome.failure is None
    assert [result.sql_table_name for result in outcome.tables] == [
        "figures_xlsx_quarterly_department_figures"
    ]
    assert outcome.tables[0].row_count == 3
    assert {sheet for sheet, _ in outcome.skipped} == {"Raw", "Blank"}

    # The folder is still the folder: its status is the documents', and the
    # sandbox database is the one C3 requires, not Askwell's own.
    status, database = await _source_row(factory, source_id)
    assert status == "ready"
    assert database is not None
    tables = _sandbox_tables(table_settings, database)
    assert tables == {"figures_xlsx_quarterly_department_figures": "askwell workbook: figures.xlsx"}
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        headcount = conn.execute(
            "SELECT headcount FROM figures_xlsx_quarterly_department_figures "
            "WHERE department = 'Logistics'"
        ).fetchone()
    assert headcount == (19,)

    async with factory() as session:
        notes = (
            await session.execute(
                text(
                    "SELECT table_name, column_name, description FROM schema_notes "
                    "WHERE source_id = :id AND superseded_by IS NULL"
                ),
                {"id": source_id},
            )
        ).all()
        # Nothing asked: a folder's clarifications are its documents' (#851).
        clarifications = (
            await session.execute(
                text("SELECT count(*) FROM clarifications WHERE source_id = :id"),
                {"id": source_id},
            )
        ).scalar_one()
    assert clarifications == 0
    assert {table for table, _, _ in notes} == {"figures.xlsx:Quarterly Department Figures"}
    assert {column for _, column, _ in notes} >= {None, "Headcount", "Q1 Revenue"}
    table_note = next(description for _, column, description in notes if column is None)
    assert "figures_xlsx_quarterly_department_figures" in table_note

    [loaded] = await _audit(factory, WORKBOOK_TABLES_LOADED)
    assert loaded["source_id"] == str(source_id)
    assert loaded["document_id"] == str(document_id)
    assert loaded["tables"] == ["figures_xlsx_quarterly_department_figures"]
    assert {entry["sheet"] for entry in loaded["skipped"]} == {"Raw", "Blank"}  # type: ignore[union-attr]


async def test_two_sheets_with_the_same_headers_become_two_tables(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "regions.xlsx"
    _write_workbook(workbook, {"North": _FIGURES, "South": _FIGURES[:2]})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)

    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert [(r.sql_table_name, r.row_count) for r in outcome.tables] == [
        ("regions_xlsx_north", 3),
        ("regions_xlsx_south", 1),
    ]


async def test_reloading_one_workbook_replaces_its_own_tables_and_leaves_anothers_alone(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    # Two names that normalise to the same identifier.
    first, second = root / "a b.xlsx", root / "a_b.xlsx"
    _write_workbook(first, {"Data": _FIGURES})
    _write_workbook(second, {"Data": _FIGURES[:2]})
    source_id = await _folder_source(factory, root)
    first_id = await _workbook_document(factory, source_id, first)
    second_id = await _workbook_document(factory, source_id, second)

    await load_workbook_tables(factory, table_settings, source_id, first_id)
    await load_workbook_tables(factory, table_settings, source_id, second_id)
    # The file changed on disk and was ingested again.
    _write_workbook(first, {"Renamed": _FIGURES})
    await load_workbook_tables(factory, table_settings, source_id, first_id)

    _, database = await _source_row(factory, source_id)
    assert database is not None
    assert _sandbox_tables(table_settings, database) == {
        "a_b_xlsx_data_2": "askwell workbook: a_b.xlsx",
        "a_b_xlsx_renamed": "askwell workbook: a b.xlsx",
    }
    async with factory() as session:
        tables = set(
            (
                await session.execute(
                    text(
                        "SELECT DISTINCT table_name FROM schema_notes "
                        "WHERE source_id = :id AND superseded_by IS NULL"
                    ),
                    {"id": source_id},
                )
            ).scalars()
        )
    assert tables == {"a b.xlsx:Renamed", "a_b.xlsx:Data"}


async def test_a_sheet_over_the_size_cap_fails_without_failing_the_document(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "big.xlsx"
    _write_workbook(workbook, {"Big": [["name", "amount"]] + [[f"row{i}", i] for i in range(500)]})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024)
        await session.commit()

    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert outcome.tables == []
    assert outcome.failure is not None
    assert "size cap" in outcome.failure
    # Nothing half-loaded is left to answer from, and the folder is untouched.
    status, database = await _source_row(factory, source_id)
    assert status == "ready"
    assert database is None
    [failed] = await _audit(factory, WORKBOOK_TABLES_FAILED)
    assert failed["source_id"] == str(source_id)
    assert failed["cap"] == "size"


async def test_a_sheet_over_the_time_cap_is_stopped_without_failing_the_document(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """The time cap still bounds a sheet's load, checked between batches
    (`M11-FIX-ING-234`): the load stops, cleans up after itself, and records
    why. The cap is set far below one batch's insert time so the first check
    trips it."""
    from askwell.dump_import import set_dump_time_cap_seconds

    root = tmp_path / "work"
    workbook = root / "large.xlsx"
    _write_workbook(
        workbook, {"Ledger": [["name", "amount"]] + [[f"row{i}", i] for i in range(5_000)]}
    )
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    async with factory() as session:
        await set_dump_time_cap_seconds(session, 0.001)
        await session.commit()

    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert outcome.tables == []
    assert outcome.failure is not None
    assert "time cap" in outcome.failure
    status, database = await _source_row(factory, source_id)
    assert status == "ready"
    assert database is None
    [failed] = await _audit(factory, WORKBOOK_TABLES_FAILED)
    assert failed["source_id"] == str(source_id)
    assert failed["cap"] == "time"


async def test_a_100k_row_sheet_loads_under_the_default_caps(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """`M11-FIX-ING-234`, #862: a sheet this size used to run past the
    default 600 s time cap, one autocommitted insert per row. Nothing here
    lowers or raises a cap."""
    root = tmp_path / "work"
    workbook = root / "large.xlsx"
    _write_workbook(
        workbook, {"Ledger": [["name", "amount"]] + [[f"row{i}", i] for i in range(100_000)]}
    )
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)

    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert outcome.failure is None
    [result] = outcome.tables
    assert result.row_count == 100_000
    assert result.failed_rows == []
    _, database = await _source_row(factory, source_id)
    assert database is not None
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        counted = conn.execute(
            f"SELECT count(*), sum(amount) FROM {result.sql_table_name}"
        ).fetchone()
    assert counted == (100_000, sum(range(100_000)))


async def test_deleting_a_workbook_takes_its_tables_out_of_every_answer_then_drops_them(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """#852. The notes go with the delete — they are all SQL generation and
    routing ever see — and the tables go at the next sweep. Another
    workbook in the same folder is untouched."""
    from askwell.sources import delete_document

    root = tmp_path / "work"
    gone, kept = root / "gone.xlsx", root / "kept.xlsx"
    _write_workbook(gone, {"Data": _FIGURES})
    _write_workbook(kept, {"Data": _FIGURES})
    source_id = await _folder_source(factory, root)
    gone_id = await _workbook_document(factory, source_id, gone)
    kept_id = await _workbook_document(factory, source_id, kept)
    await load_workbook_tables(factory, table_settings, source_id, gone_id)
    await load_workbook_tables(factory, table_settings, source_id, kept_id)
    async with factory() as session:
        await session.execute(
            text(
                "INSERT INTO schema_notes (id, source_id, table_name, column_name, description, "
                "origin, confidence) VALUES (:id, :source_id, 'gone.xlsx:Data', 'Headcount', "
                "'People on the payroll.', 'user', 1.0)"
            ),
            {"id": uuid.uuid4(), "source_id": source_id},
        )
        assert await delete_document(session, gone_id, None, 0.6, table_settings)
        await session.commit()

        notes = (
            await session.execute(
                text(
                    "SELECT table_name, origin, stale FROM schema_notes "
                    "WHERE source_id = :id AND left(table_name, 10) = 'gone.xlsx:'"
                ),
                {"id": source_id},
            )
        ).all()
    assert notes == [("gone.xlsx:Data", "user", True)]

    assert await sweep_workbook_tables(factory, table_settings) == 1
    _, database = await _source_row(factory, source_id)
    assert database is not None
    assert _sandbox_tables(table_settings, database) == {
        "kept_xlsx_data": "askwell workbook: kept.xlsx"
    }
    # Nothing left to drop, and the owner's grant is not reopened for it.
    assert await sweep_workbook_tables(factory, table_settings) == 0


async def test_deleting_an_older_version_leaves_the_newer_ones_tables(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell.table_load import forget_deleted_workbook

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Data": _FIGURES})
    source_id = await _folder_source(factory, root)
    older = await _workbook_document(factory, source_id, workbook)
    newer = await _workbook_document(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, newer)

    async with factory() as session:
        await session.execute(
            text("UPDATE documents SET deleted_at = now(), status = 'deleted' WHERE id = :id"),
            {"id": older},
        )
        await forget_deleted_workbook(session, older)
        await session.commit()
        remaining = (
            await session.execute(
                text("SELECT count(*) FROM schema_notes WHERE source_id = :id AND NOT stale"),
                {"id": source_id},
            )
        ).scalar_one()
    assert remaining > 0
    assert await sweep_workbook_tables(factory, table_settings) == 0


async def test_a_sheets_inferred_notes_stay_out_of_a_document_answer_and_in_sql_generation(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """#857. The same workbook is also passages, which carry the same numbers
    with a citation. A note Askwell inferred about a sheet column ("Loaded
    as column `q1_revenue`") is for writing SQL; in a document answer it
    competes with the passage and gets cited in its place. What a person
    said about a sheet is memory, and still applies."""
    from askwell.memory import retrieve_relevant_facts

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, document_id)
    question = "What was the Q1 revenue of the Research department, per the figures?"

    async with factory() as session:
        for_sql = await retrieve_relevant_facts(session, question=question, source_id=source_id)
        for_documents = await retrieve_relevant_facts(
            session, question=question, include_inferred_sheet_notes=False
        )
    assert {note.origin for note in for_sql.notes} == {"inferred"}
    assert for_documents.notes == []

    async with factory() as session:
        await session.execute(
            text(
                "INSERT INTO schema_notes (id, source_id, table_name, column_name, description, "
                "origin, confidence) VALUES (:id, :source_id, 'figures.xlsx:Figures', "
                "'Q1 Revenue', 'Q1 revenue is in US dollars.', 'user', 1.0)"
            ),
            {"id": uuid.uuid4(), "source_id": source_id},
        )
        await session.commit()
        for_documents = await retrieve_relevant_facts(
            session, question=question, include_inferred_sheet_notes=False
        )
    assert [(note.origin, note.description) for note in for_documents.notes] == [
        ("user", "Q1 revenue is in US dollars.")
    ]


# --- a workbook's sheets ask about what they could not type (`M11-FIX-ING-228`) ---


# Both parts of every `Placed` value are `<= 12`, so no value settles DD/MM
# against MM/DD and the column loads as `text` until someone says which.
_ORDERS: list[list[object]] = [
    ["Customer", "Placed", "Units"],
    ["Anna", "03/04/2025", 4],
    ["Ben", "05/06/2025", 7],
    ["Cara", "11/12/2025", 2],
]


async def _document_with_an_abbreviation(
    factory: async_sessionmaker[AsyncSession], source_id: uuid.UUID, root: Path
) -> None:
    async with factory() as session:
        document_id = (
            await session.execute(
                text(
                    "INSERT INTO documents (source_id, filename, path, mime, sha256, status) "
                    "VALUES (:source_id, 'tender.pdf', :path, 'application/pdf', :sha256, "
                    "'ready') RETURNING id"
                ),
                {
                    "source_id": source_id,
                    "path": str(root / "tender.pdf"),
                    "sha256": hashlib.sha256(os.urandom(8)).hexdigest(),
                },
            )
        ).scalar_one()
        for ordinal, content in enumerate(
            ("The RFQ closes Friday.", "Submit the RFQ to procurement.")
        ):
            await session.execute(
                text(
                    "INSERT INTO chunks (id, document_id, ordinal, content) "
                    "VALUES (:id, :document_id, :ordinal, :content)"
                ),
                {
                    "id": uuid.uuid4(),
                    "document_id": document_id,
                    "ordinal": ordinal,
                    "content": content,
                },
            )
        await session.commit()


async def _raise_document_questions(
    factory: async_sessionmaker[AsyncSession], settings: Settings, source_id: uuid.UUID
) -> None:
    from askwell.clarify import raise_candidates

    async with factory() as session:
        await raise_candidates(session, source_id, 0.6, settings)
        await session.commit()


async def _questions(
    factory: async_sessionmaker[AsyncSession], source_id: uuid.UUID
) -> list[tuple[str, str, str | None]]:
    async with factory() as session:
        rows = await session.execute(
            text(
                "SELECT subject, question, evidence->>'workbook' FROM clarifications "
                "WHERE source_id = :id ORDER BY subject"
            ),
            {"id": source_id},
        )
        return [(row[0], row[1], row[2]) for row in rows]


def _column_type(settings: Settings, database: str, table: str, column: str) -> str:
    admin_url = settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        row = conn.execute(
            "SELECT data_type FROM information_schema.columns "
            "WHERE table_name = %s AND column_name = %s",
            (table, column),
        ).fetchone()
    assert row is not None
    return str(row[0])


async def test_a_folder_with_a_document_and_a_workbook_asks_both_kinds_of_question(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "orders.xlsx"
    _write_workbook(workbook, {"North": _ORDERS})
    source_id = await _folder_source(factory, root)
    await _document_with_an_abbreviation(factory, source_id, root)
    document_id = await _workbook_document(factory, source_id, workbook)

    # The order ingestion runs them in: a workbook's sheets load while it is
    # indexed, the folder is scanned once nothing is left outstanding.
    await load_workbook_tables(factory, table_settings, source_id, document_id)
    await _raise_document_questions(factory, table_settings, source_id)

    questions = await _questions(factory, source_id)
    assert [(subject, workbook) for subject, _question, workbook in questions] == [
        ("Placed", "orders.xlsx"),
        ("RFQ", None),
    ]
    # Reads like a CSV's, and says which workbook and sheet it is about.
    assert questions[0][1].startswith(
        "*Placed* in *orders.xlsx*, sheet *North*, looks like a date in DD/MM/YYYY or MM/DD/YYYY"
    )
    assert await _audit(factory, "table_clarification_raised") == [
        {
            "source_id": str(source_id),
            "trigger": "date_format",
            "subject": "Placed",
            "rank": 1,
            "workbook": "orders.xlsx",
        }
    ]


async def test_a_folder_asked_about_its_documents_before_this_change_still_asks_about_a_sheet(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "orders.xlsx"
    _write_workbook(workbook, {"North": _ORDERS})
    source_id = await _folder_source(factory, root)
    await _document_with_an_abbreviation(factory, source_id, root)
    await _raise_document_questions(factory, table_settings, source_id)
    document_id = await _workbook_document(factory, source_id, workbook)

    await load_workbook_tables(factory, table_settings, source_id, document_id)
    # A later scan of the folder asks nothing new about its documents.
    await _raise_document_questions(factory, table_settings, source_id)

    assert [(s, w) for s, _q, w in await _questions(factory, source_id)] == [
        ("Placed", "orders.xlsx"),
        ("RFQ", None),
    ]


async def test_a_workbooks_questions_are_capped_with_the_folders_pending_ones(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell.clarify import set_clarification_cap

    root = tmp_path / "work"
    workbook = root / "orders.xlsx"
    _write_workbook(
        workbook,
        {
            "North": _ORDERS,
            "South": [["Customer", "Shipped", "Units"], *_ORDERS[1:]],
        },
    )
    source_id = await _folder_source(factory, root)
    await _document_with_an_abbreviation(factory, source_id, root)
    async with factory() as session:
        await set_clarification_cap(session, 2)
        await session.commit()
    await _raise_document_questions(factory, table_settings, source_id)
    document_id = await _workbook_document(factory, source_id, workbook)

    await load_workbook_tables(factory, table_settings, source_id, document_id)

    # One slot was taken by the folder's pending document question.
    assert len(await _questions(factory, source_id)) == 2
    capped = await _audit(factory, "clarification_capped")
    assert len(capped) == 1
    assert capped[0]["workbook"] == "orders.xlsx"


async def test_answering_a_sheets_date_question_turns_the_column_into_a_date(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell import reapply

    root = tmp_path / "work"
    orders, other = root / "orders.xlsx", root / "other.xlsx"
    _write_workbook(orders, {"North": _ORDERS})
    _write_workbook(other, {"North": _ORDERS})
    source_id = await _folder_source(factory, root)
    orders_id = await _workbook_document(factory, source_id, orders)
    other_id = await _workbook_document(factory, source_id, other)
    await load_workbook_tables(factory, table_settings, source_id, orders_id)
    await load_workbook_tables(factory, table_settings, source_id, other_id)
    _status, database = await _source_row(factory, source_id)
    assert database is not None
    assert _column_type(table_settings, database, "orders_xlsx_north", "placed") == "text"

    async with factory() as session:
        clarification_id, options = (
            await session.execute(
                text(
                    "SELECT id, options FROM clarifications WHERE source_id = :id "
                    "AND evidence->>'workbook' = 'orders.xlsx'"
                ),
                {"id": source_id},
            )
        ).one()
        outcome = await answer_clarification(session, clarification_id, options[0])
        await session.commit()
    assert outcome.reapply_job_id is not None

    await reapply.run_job(factory, table_settings, outcome.reapply_job_id)

    assert _column_type(table_settings, database, "orders_xlsx_north", "placed") == "date"
    # The other workbook's own question is still open, so its column is not.
    assert _column_type(table_settings, database, "other_xlsx_north", "placed") == "text"
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    with psycopg.connect(admin_url.rsplit("/", 1)[0] + f"/{database}", autocommit=True) as conn:
        placed = conn.execute(
            "SELECT customer, placed FROM orders_xlsx_north ORDER BY customer"
        ).fetchall()
        comment = conn.execute(
            "SELECT obj_description('orders_xlsx_north'::regclass, 'pg_class')"
        ).fetchone()
    # Day first: `03/04/2025` is 3 April.
    assert [(name, value.isoformat()) for name, value in placed] == [
        ("Anna", "2025-04-03"),
        ("Ben", "2025-06-05"),
        ("Cara", "2025-12-11"),
    ]
    assert comment == ("askwell workbook: orders.xlsx",)
    assert await _audit(factory, "table_column_reloaded") == [
        {
            "source_id": str(source_id),
            "workbook": "orders.xlsx",
            "tables": ["orders_xlsx_north"],
            "failed_rows": 0,
        }
    ]


async def test_a_workbook_re_ingested_after_its_question_was_answered_keeps_the_answer(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell import reapply

    root = tmp_path / "work"
    workbook = root / "orders.xlsx"
    _write_workbook(workbook, {"North": _ORDERS})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, document_id)
    async with factory() as session:
        clarification_id, options = (
            await session.execute(
                text("SELECT id, options FROM clarifications WHERE source_id = :id"),
                {"id": source_id},
            )
        ).one()
        outcome = await answer_clarification(session, clarification_id, options[1])
        await session.commit()
    assert outcome.reapply_job_id is not None
    await reapply.run_job(factory, table_settings, outcome.reapply_job_id)

    # The file changes and is indexed again, as a new version of itself.
    _write_workbook(workbook, {"North": [*_ORDERS, ["Dev", "07/08/2025", 1]]})
    newer_id = await _workbook_document(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, newer_id)

    _status, database = await _source_row(factory, source_id)
    assert database is not None
    assert _column_type(table_settings, database, "orders_xlsx_north", "placed") == "date"
    # Asked once: the answered question is the only one.
    assert len(await _questions(factory, source_id)) == 1
    async with factory() as session:
        notes = (
            await session.execute(
                text(
                    "SELECT origin, description FROM schema_notes WHERE source_id = :id "
                    "AND column_name = 'Placed' AND superseded_by IS NULL"
                ),
                {"id": source_id},
            )
        ).all()
    # The person's answer is the column's note; no fresh guess sits beside it.
    assert notes == [("user", options[1])]


# --- the library says when a sheet was not loaded (`M11-FIX-UI-229`) ---------


async def _sheet_outcome(
    factory: async_sessionmaker[AsyncSession], document_id: uuid.UUID
) -> tuple[object, object]:
    async with factory() as session:
        row = (
            await session.execute(
                text("SELECT sheet_load_failure, sheets_skipped FROM documents WHERE id = :id"),
                {"id": document_id},
            )
        ).one()
    return row[0], row[1]


async def _refresh(
    factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    source_id: uuid.UUID,
    document_id: uuid.UUID,
) -> tuple[str, str | None]:
    """The document finished indexing: what `ingest._finish` does next."""
    async with factory() as session:
        await session.execute(
            text("UPDATE documents SET status = 'ready' WHERE id = :id"), {"id": document_id}
        )
        await ingest.refresh_source(session, source_id, settings.ocr_confidence_threshold, settings)
        await session.commit()
        row = (
            await session.execute(
                text("SELECT status, last_error FROM sources WHERE id = :id"), {"id": source_id}
            )
        ).one()
    return row[0], row[1]


async def test_a_skipped_sheet_is_a_note_on_its_document_not_attention(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES, "Raw": [[1, 2], [3, 4], [5, 6]]})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)

    await load_workbook_tables(factory, table_settings, source_id, document_id)

    failure, skipped = await _sheet_outcome(factory, document_id)
    assert failure is None
    assert isinstance(skipped, list)
    [note] = skipped
    assert note["sheet"] == "Raw"
    assert note["reason"].startswith("no recognisable header row")

    # Nothing failed, so the folder is ready — and the library still hears it.
    assert await _refresh(factory, table_settings, source_id, document_id) == ("ready", None)
    async with factory() as session:
        snapshot = await ingest.snapshot(session, table_settings)
    [entry] = snapshot["sheet_notes"]
    assert entry["document_id"] == str(document_id)
    assert entry["filename"] == "figures.xlsx"
    assert entry["failure"] is None
    assert [item["sheet"] for item in entry["skipped"]] == ["Raw"]
    [source] = snapshot["sources"]
    assert source["sheets_failed"] == 0


async def test_a_failed_sheet_load_puts_the_folder_in_attention_naming_workbook_sheet_and_reason(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "big.xlsx"
    _write_workbook(workbook, {"Big": [["name", "amount"]] + [[f"row{i}", i] for i in range(500)]})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024)
        await session.commit()

    await load_workbook_tables(factory, table_settings, source_id, document_id)

    failure, _ = await _sheet_outcome(factory, document_id)
    assert isinstance(failure, dict)
    assert failure["sheet"] == "Big"
    assert "size cap" in failure["reason"]

    status, reason = await _refresh(factory, table_settings, source_id, document_id)
    assert status == "attention"
    assert reason is not None
    assert reason.startswith("big.xlsx: the sheet Big was not loaded as a table. Load aborted")
    assert "size cap" in reason

    async with factory() as session:
        snapshot = await ingest.snapshot(session, table_settings)
    [source] = snapshot["sources"]
    assert source["sheets_failed"] == 1
    [entry] = snapshot["sheet_notes"]
    assert entry["failure"]["sheet"] == "Big"


async def test_a_workbook_re_ingested_successfully_clears_the_note(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """The ticket's edge case: the cap is raised, the workbook indexed again,
    and neither the attention nor the note survives it."""
    root = tmp_path / "work"
    workbook = root / "big.xlsx"
    _write_workbook(workbook, {"Big": [["name", "amount"]] + [[f"row{i}", i] for i in range(500)]})
    source_id = await _folder_source(factory, root)
    document_id = await _workbook_document(factory, source_id, workbook)
    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024)
        await session.commit()
    await load_workbook_tables(factory, table_settings, source_id, document_id)
    assert (await _refresh(factory, table_settings, source_id, document_id))[0] == "attention"

    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024 * 1024 * 1024)
        await session.commit()
    outcome = await load_workbook_tables(factory, table_settings, source_id, document_id)

    assert outcome.failure is None
    assert await _sheet_outcome(factory, document_id) == (None, None)
    assert await _refresh(factory, table_settings, source_id, document_id) == ("ready", None)
    async with factory() as session:
        snapshot = await ingest.snapshot(session, table_settings)
    assert snapshot["sheet_notes"] == []


async def test_a_newer_version_of_the_workbook_hides_the_older_versions_note(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Raw": [[1, 2], [3, 4], [5, 6]]})
    source_id = await _folder_source(factory, root)
    older = await _workbook_document(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, older)

    _write_workbook(workbook, {"Figures": _FIGURES})
    newer = await _workbook_document(factory, source_id, workbook)
    async with factory() as session:
        await session.execute(
            text("UPDATE documents SET superseded_by = :newer WHERE id = :older"),
            {"newer": newer, "older": older},
        )
        await session.commit()
    await load_workbook_tables(factory, table_settings, source_id, newer)

    async with factory() as session:
        snapshot = await ingest.snapshot(session, table_settings)
    assert snapshot["sheet_notes"] == []


# --- workbooks indexed before their sheets became tables (`M11-FIX-ING-230`) -


async def _indexed_before_tables(
    factory: async_sessionmaker[AsyncSession], source_id: uuid.UUID, path: Path
) -> uuid.UUID:
    """A workbook `ready` before `0.9.12`: indexed, with no load decision."""
    document_id = await _workbook_document(factory, source_id, path)
    async with factory() as session:
        await session.execute(
            text("UPDATE documents SET status = 'ready' WHERE id = :id"), {"id": document_id}
        )
        await session.commit()
    return document_id


async def test_a_workbook_indexed_before_its_sheets_became_tables_gets_them_at_start(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell.table_load import backfill_workbook_tables

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    document_id = await _indexed_before_tables(factory, source_id, workbook)

    assert await backfill_workbook_tables(factory, table_settings) == 1

    status, database = await _source_row(factory, source_id)
    assert status == "ready"
    assert database is not None
    assert _sandbox_tables(table_settings, database) == {
        "figures_xlsx_figures": "askwell workbook: figures.xlsx"
    }
    [loaded] = await _audit(factory, WORKBOOK_TABLES_LOADED)
    assert loaded["document_id"] == str(document_id)
    async with factory() as session:
        # No re-embedding: the document is left exactly as indexed.
        document_status = (
            await session.execute(
                text("SELECT status FROM documents WHERE id = :id"), {"id": document_id}
            )
        ).scalar_one()
    assert document_status == "ready"

    # Idempotent: the next start finds nothing to do.
    assert await backfill_workbook_tables(factory, table_settings) == 0
    assert len(await _audit(factory, WORKBOOK_TABLES_LOADED)) == 1


async def test_a_workbook_whose_sheets_were_all_skipped_is_not_rescanned_at_every_start(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell.table_load import backfill_workbook_tables

    root = tmp_path / "work"
    workbook = root / "raw.xlsx"
    _write_workbook(workbook, {"Raw": [[1, 2], [3, 4], [5, 6]]})
    source_id = await _folder_source(factory, root)
    document_id = await _indexed_before_tables(factory, source_id, workbook)

    assert await backfill_workbook_tables(factory, table_settings) == 1
    # Nothing loadable, so no database — but the decision is the marker.
    assert (await _source_row(factory, source_id))[1] is None
    [loaded] = await _audit(factory, WORKBOOK_TABLES_LOADED)
    assert loaded["tables"] == []
    failure, skipped = await _sheet_outcome(factory, document_id)
    assert failure is None
    assert [item["sheet"] for item in skipped] == ["Raw"]  # type: ignore[union-attr]

    assert await backfill_workbook_tables(factory, table_settings) == 0
    assert len(await _audit(factory, WORKBOOK_TABLES_LOADED)) == 1


async def test_the_backfill_leaves_failed_deleted_superseded_and_unfinished_workbooks_alone(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    from askwell.table_load import backfill_workbook_tables

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    failed = await _indexed_before_tables(factory, source_id, workbook)
    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024)
        await session.commit()
    assert (await load_workbook_tables(factory, table_settings, source_id, failed)).failure
    async with factory() as session:
        await set_dump_size_cap_bytes(session, 1024 * 1024 * 1024)
        await session.commit()
    deleted = await _indexed_before_tables(factory, source_id, workbook)
    superseded = await _indexed_before_tables(factory, source_id, workbook)
    unfinished = await _workbook_document(factory, source_id, workbook)  # still indexing
    async with factory() as session:
        await session.execute(
            text("UPDATE documents SET deleted_at = now(), status = 'deleted' WHERE id = :id"),
            {"id": deleted},
        )
        await session.execute(
            text("UPDATE documents SET superseded_by = :newer WHERE id = :older"),
            {"newer": unfinished, "older": superseded},
        )
        await session.commit()

    assert await backfill_workbook_tables(factory, table_settings) == 0
    assert await _audit(factory, WORKBOOK_TABLES_LOADED) == []
    assert len(await _audit(factory, WORKBOOK_TABLES_FAILED)) == 1


async def test_a_workbook_deleted_between_the_scan_and_the_load_is_not_loaded(
    factory: async_sessionmaker[AsyncSession],
    table_settings: Settings,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from askwell import table_load

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    document_id = await _indexed_before_tables(factory, source_id, workbook)

    scan = table_load._workbooks_without_tables

    async def scan_then_delete(
        session: AsyncSession,
    ) -> list[tuple[uuid.UUID, uuid.UUID]]:
        found = await scan(session)
        async with factory() as other:
            await other.execute(
                text("UPDATE documents SET deleted_at = now(), status = 'deleted' WHERE id = :id"),
                {"id": document_id},
            )
            await other.commit()
        return found

    monkeypatch.setattr(table_load, "_workbooks_without_tables", scan_then_delete)

    assert await table_load.backfill_workbook_tables(factory, table_settings) == 0
    # Nothing created, nothing recorded: a deleted workbook is not a failure.
    assert (await _source_row(factory, source_id))[1] is None
    assert await _audit(factory, WORKBOOK_TABLES_LOADED) == []
    assert await _audit(factory, WORKBOOK_TABLES_FAILED) == []


async def test_a_sandbox_that_is_not_up_at_worker_start_defers_the_backfill_not_fails_it(
    factory: async_sessionmaker[AsyncSession],
    table_settings: Settings,
    tmp_path: Path,
    database_url: str,
) -> None:
    from askwell import worker
    from askwell.table_load import backfill_workbook_tables

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    await _indexed_before_tables(factory, source_id, workbook)
    down = table_settings.model_copy(
        update={"sandbox_database_url": SecretStr("postgresql://x:x@127.0.0.1:1/askwell_sandbox")}
    )

    assert await worker._reclaim_sandbox_orphans(factory, down) == ()
    assert await _audit(factory, WORKBOOK_TABLES_FAILED) == []
    assert await _audit(factory, WORKBOOK_TABLES_LOADED) == []

    # The next start, with the sandbox up, does it.
    assert await backfill_workbook_tables(factory, table_settings) == 1
    assert len(await _audit(factory, WORKBOOK_TABLES_LOADED)) == 1


async def test_a_folder_whose_sandbox_database_is_gone_gets_its_tables_back(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """#895. A workbook marked loaded into a database that no longer exists
    is eligible again, once, and the folder gets a new database."""
    from askwell import sandbox
    from askwell.table_load import (
        SANDBOX_DATABASE_LOST,
        backfill_workbook_tables,
        forget_lost_databases,
    )

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    document_id = await _indexed_before_tables(factory, source_id, workbook)
    await load_workbook_tables(factory, table_settings, source_id, document_id)
    _, lost = await _source_row(factory, source_id)
    assert lost is not None
    admin_url = table_settings.sandbox_database_url.get_secret_value()
    async with factory() as session:
        await sandbox.drop_database(session, admin_url, lost, reason="test")
        await session.commit()

    assert await forget_lost_databases(factory, table_settings) == [source_id]
    status, database = await _source_row(factory, source_id)
    assert (status, database) == ("ready", None)
    [recorded] = await _audit(factory, SANDBOX_DATABASE_LOST)
    assert recorded == {"source_id": str(source_id), "database": lost}
    async with factory() as session:
        inferred = (
            await session.execute(
                text(
                    "SELECT count(*) FROM schema_notes "
                    "WHERE source_id = :id AND origin = 'inferred'"
                ),
                {"id": source_id},
            )
        ).scalar_one()
    # No note points SQL generation at a table that is not there.
    assert inferred == 0

    assert await backfill_workbook_tables(factory, table_settings) == 1
    _, database = await _source_row(factory, source_id)
    assert database is not None and database != lost
    assert _sandbox_tables(table_settings, database) == {
        "figures_xlsx_figures": "askwell workbook: figures.xlsx"
    }

    # Once: nothing more is lost, and nothing is loaded again.
    assert await forget_lost_databases(factory, table_settings) == []
    assert await backfill_workbook_tables(factory, table_settings) == 0


async def test_a_workbook_whose_file_cannot_be_read_at_start_is_deferred_not_failed(
    factory: async_sessionmaker[AsyncSession], table_settings: Settings, tmp_path: Path
) -> None:
    """A folder outside the mount, or a drive not plugged in: the next start
    tries again, rather than a failure stopping it for good."""
    from askwell.table_load import backfill_workbook_tables

    root = tmp_path / "work"
    workbook = root / "figures.xlsx"
    _write_workbook(workbook, {"Figures": _FIGURES})
    source_id = await _folder_source(factory, root)
    await _indexed_before_tables(factory, source_id, workbook)
    hidden = root.with_name("elsewhere")
    root.rename(hidden)

    assert await backfill_workbook_tables(factory, table_settings) == 0
    assert await _audit(factory, WORKBOOK_TABLES_FAILED) == []
    assert await _audit(factory, WORKBOOK_TABLES_LOADED) == []

    hidden.rename(root)
    assert await backfill_workbook_tables(factory, table_settings) == 1
    assert len(await _audit(factory, WORKBOOK_TABLES_LOADED)) == 1
